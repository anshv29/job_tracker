"""Appends today's submitted jobs to a career-tracker Excel file, if one exists.

There is no tracker file in this repo today, so this is a no-op until/unless
you add one - it never errors, it just logs and returns.
"""
import os

TRACKER_PATH = "career_tracker.xlsx"
COLUMNS = ["Date Applied", "Company", "Title", "ATS", "Location", "URL"]


def append_submitted_jobs(rows, path=TRACKER_PATH):
    if not os.path.exists(path):
        print(f"{path} doesn't exist in this repo, skipping the Excel export.")
        return

    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    if ws.max_row == 0 or ws.cell(1, 1).value is None:
        ws.append(COLUMNS)

    for row in rows:
        ws.append([
            row["applied_at"].strftime("%Y-%m-%d") if row.get("applied_at") else "",
            row["company"], row["title"], row["ats"], row.get("location") or "", row["url"],
        ])
    wb.save(path)
    print(f"Appended {len(rows)} row(s) to {path}")
